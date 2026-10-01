from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "docx" / "The-Dialogue-Platform-Fritt-Ord-Concept-Note.docx"
LOGO = ROOT / "public" / "assets" / "logo.png"

NAVY = "123E5E"
ORANGE = "F4A228"
INK = "202A33"
MUTED = "5B6873"
PALE = "F7F3EA"
WHITE = "FFFFFF"
LINE = "D9D2C5"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=110, bottom=80, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=LINE, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_text(paragraph, text, *, size=9.0, bold=False, color=INK, font="Aptos"):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_label(cell, text, *, light=False):
    p = cell.add_paragraph() if cell.paragraphs[0].text else cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    add_text(p, text.upper(), size=7.4, bold=True, color=WHITE if light else ORANGE)
    return p


def add_body(cell, text, *, size=8.6, color=INK, bold=False, after=3):
    p = cell.add_paragraph() if cell.paragraphs[0].text else cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.03
    add_text(p, text, size=size, bold=bold, color=color)
    return p


def add_bullet(cell, text, *, color=INK, size=8.4, after=1.5):
    p = cell.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.35)
    p.paragraph_format.first_line_indent = Cm(-0.22)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    add_text(p, text, size=size, color=color)
    return p


def set_table_cell_width(cell, width_cm):
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(0.85)
    section.bottom_margin = Cm(0.75)
    section.left_margin = Cm(1.0)
    section.right_margin = Cm(1.0)
    section.header_distance = Cm(0.3)
    section.footer_distance = Cm(0.3)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9)
    styles["Normal"].font.color.rgb = RGBColor.from_string(INK)
    styles["Normal"].paragraph_format.space_after = Pt(3)

    header = doc.add_table(rows=1, cols=2)
    header.alignment = WD_TABLE_ALIGNMENT.CENTER
    header.autofit = False
    left, right = header.rows[0].cells
    set_table_cell_width(left, 5.0)
    set_table_cell_width(right, 14.0)
    for cell in (left, right):
        set_cell_margins(cell, 0, 0, 0, 0)
    if LOGO.exists():
        p = left.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.add_run().add_picture(str(LOGO), width=Cm(3.1))
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    add_text(p, "PROJECT CONCEPT NOTE", size=8, bold=True, color=ORANGE)
    p = right.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    add_text(p, "Prepared for Fritt Ord", size=8.5, bold=True, color=NAVY)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(7)
    title.paragraph_format.space_after = Pt(1)
    add_text(title, "Dialogue Rooms Norway", size=21, bold=True, color=NAVY, font="Aptos Display")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(7)
    add_text(
        subtitle,
        "Four public meetings for freedom of expression, trust and democratic participation",
        size=10.4,
        bold=True,
        color=ORANGE,
    )

    facts = doc.add_table(rows=1, cols=4)
    facts.alignment = WD_TABLE_ALIGNMENT.CENTER
    facts.autofit = False
    fact_data = [
        ("Applicant", "The Dialogue Platform"),
        ("Proposed period", "February-June 2027"),
        ("Locations", "Lillestrom and Oslo"),
        ("Fritt Ord request", "NOK 100,000"),
    ]
    for cell, (label, value) in zip(facts.rows[0].cells, fact_data):
        set_table_cell_width(cell, 4.75)
        set_cell_margins(cell, 95, 105, 95, 105)
        shade(cell, NAVY)
        add_label(cell, label, light=True)
        add_body(cell, value, size=8.5, color=WHITE, bold=True, after=0)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.space_before = Pt(0)

    columns = doc.add_table(rows=1, cols=2)
    columns.alignment = WD_TABLE_ALIGNMENT.CENTER
    columns.autofit = False
    left_col, right_col = columns.rows[0].cells
    set_table_cell_width(left_col, 11.2)
    set_table_cell_width(right_col, 7.8)
    for cell in (left_col, right_col):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_margins(cell, 0, 0, 0, 0)

    add_label(left_col, "Purpose and public need")
    add_body(
        left_col,
        "Public debate is increasingly shaped by polarization, distrust and unequal access to the spaces where opinions are formed. Minority, diaspora, youth and women voices are often discussed without being fully included. Dialogue Rooms Norway will create open, carefully moderated physical meetings where disagreement can be expressed safely, tested against other experiences and turned into democratic participation.",
        size=8.7,
        after=5,
    )

    add_label(left_col, "Activities and method")
    add_body(
        left_col,
        "The project will deliver four free public meetings, each designed for 60-100 participants. Every meeting combines a short factual framing, two or three invited perspectives, facilitated table dialogue and an open plenary. The method draws on structured exploration, exchange and reflection so participants listen actively, challenge assumptions and identify practical common ground.",
        size=8.7,
        after=3,
    )
    add_bullet(left_col, "Freedom of expression across disagreement")
    add_bullet(left_col, "Belonging, identity and democratic participation")
    add_bullet(left_col, "Media, misinformation and the language of conflict")
    add_bullet(left_col, "From polarization to practical local cooperation", after=4)

    add_label(left_col, "Public relevance and accessibility")
    add_body(
        left_col,
        "The meetings will be publicly announced and open without charge. Outreach will combine local associations, schools, municipal networks and The Dialogue Platform's digital channels. Moderation and interpretation will be adapted to participant needs. Representation across age, gender and background will be built into speaker selection, while photos or short clips will only be shared with consent.",
        size=8.6,
        after=0,
    )

    right_box = right_col.add_table(rows=1, cols=1)
    right_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    box = right_box.cell(0, 0)
    set_cell_margins(box, 120, 135, 120, 135)
    set_cell_border(box)
    shade(box, PALE)
    add_label(box, "Expected results")
    add_bullet(box, "4 moderated public dialogue meetings", size=8.3)
    add_bullet(box, "240-400 participants in total", size=8.3)
    add_bullet(box, "4 concise public event summaries", size=8.3)
    add_bullet(box, "1 final learning note with recommendations", size=8.3, after=4)

    add_label(box, "Success measures")
    add_bullet(box, "Attendance and participant diversity", size=8.3)
    add_bullet(box, "At least 75% report greater understanding of another perspective", size=8.3)
    add_bullet(box, "New collaborations or follow-up actions recorded after each meeting", size=8.3, after=4)

    add_label(box, "Provisional budget")
    budget_lines = [
        ("Facilitation and moderation", "40,000"),
        ("Venue and technical delivery", "32,000"),
        ("Guest travel and honoraria", "24,000"),
        ("Interpretation and accessibility", "20,000"),
        ("Outreach and design", "18,000"),
        ("Documentation and evaluation", "16,000"),
        ("Administration and contingency", "10,000"),
    ]
    budget = box.add_table(rows=0, cols=2)
    budget.autofit = False
    for name, amount in budget_lines:
        cells = budget.add_row().cells
        set_table_cell_width(cells[0], 4.8)
        set_table_cell_width(cells[1], 1.7)
        p = cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(0.5)
        add_text(p, name, size=7.5, color=MUTED)
        p = cells[1].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(0.5)
        add_text(p, amount, size=7.5, bold=True, color=INK)
    total = box.add_paragraph()
    total.paragraph_format.space_before = Pt(2)
    total.paragraph_format.space_after = Pt(1)
    add_text(total, "Total project cost: NOK 160,000", size=8.2, bold=True, color=NAVY)
    funding = box.add_paragraph()
    funding.paragraph_format.space_after = Pt(0)
    add_text(funding, "Fritt Ord: NOK 100,000 | Own, volunteer and partner contribution: NOK 60,000", size=7.6, bold=True, color=INK)

    org = doc.add_table(rows=1, cols=2)
    org.alignment = WD_TABLE_ALIGNMENT.CENTER
    org.autofit = False
    a, b = org.rows[0].cells
    set_table_cell_width(a, 12.2)
    set_table_cell_width(b, 6.8)
    for cell in (a, b):
        set_cell_margins(cell, 95, 110, 85, 110)
        shade(cell, NAVY)
    add_label(a, "Applicant capacity", light=True)
    add_body(
        a,
        "The Dialogue Platform is an independent, non-profit initiative based in Lillestrom. It has organized seminars, workshops and public dialogues in cooperation with community actors, with approximately 380 direct participants across 11 activities and more than 40,000 combined online views reported to date.",
        size=7.8,
        color=WHITE,
        after=0,
    )
    add_label(b, "Contact", light=True)
    add_body(b, "Omran Adam\ncontact@thedialogueplatform.com\nthedialogueplatform.com\nOrg. no. 935 674 220", size=7.8, color=WHITE, bold=True, after=0)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(
        footer,
        "DIALOGUE ROOMS NORWAY  |  ONE-PAGE PROJECT CONCEPT  |  THE DIALOGUE PLATFORM",
        size=6.8,
        bold=True,
        color=MUTED,
    )

    core = doc.core_properties
    core.title = "Dialogue Rooms Norway - Fritt Ord Concept Note"
    core.subject = "One-page project concept note for public physical dialogue meetings"
    core.author = "The Dialogue Platform"
    core.keywords = "Fritt Ord, dialogue, freedom of expression, public meetings, democracy"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
